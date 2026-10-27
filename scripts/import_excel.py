"""Read the running workbook and generate the upstream Activity data contract."""

import argparse
import hashlib
import json
import math
import os
import re
import sys
import tempfile
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

METERS_PER_MILE = 1609.344
HEADERS = [
    "Record ID",
    "Date",
    "Training type",
    "Distance (mi)",
    "Duration (min)",
    "Workout details",
    "Running steps",
    "Notes",
    "Data type",
]


def positive_number(value, label):
    if (
        isinstance(value, bool)
        or not isinstance(value, (float, int))
        or not math.isfinite(value)
        or value <= 0
    ):
        raise ValueError(f"{label} must be a positive number")
    return float(value)


def recorded_date(value, label):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            pass
    raise ValueError(f"{label} must contain a valid date")


def text(value):
    return "" if value is None else str(value).strip()


def import_workbook(path):
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "Training Log" not in workbook.sheetnames:
            raise ValueError("Missing Training Log sheet")
        rows = workbook["Training Log"].iter_rows(values_only=True)
        header = None
        for number, row in enumerate(rows, 1):
            if text(row[0]) == "Record ID":
                header = [text(v) for v in row[:9]]
                break
        if header != HEADERS:
            raise ValueError("Training Log headers must match the mile-based template")
        activities, seen = [], set()
        first_data_row = number + 1
        for number, row in enumerate(rows, first_data_row):
            row = tuple(row) + (None,) * max(0, 9 - len(row))
            if all(v is None or text(v) == "" for v in row[1:8]):
                continue
            data_type = text(row[8]).lower()
            if data_type == "example":
                continue
            if data_type != "actual":
                raise ValueError(
                    f"Row {number}: populated records require Data type Actual or Example"
                )
            record_id = text(row[0])
            if not re.fullmatch(r"R[0-9]+", record_id) or int(record_id[1:]) <= 0:
                raise ValueError(f"Row {number}: Record ID must be a positive R-number")
            run_id = int(record_id[1:])
            if run_id in seen:
                raise ValueError(f"Row {number}: duplicate Record ID {record_id}")
            seen.add(run_id)
            when = recorded_date(row[1], f"Row {number} Date")
            if when.tzinfo is not None:
                raise ValueError(
                    f"Row {number}: use a New York local date, without a timezone suffix"
                )
            distance = (
                positive_number(row[3], f"Row {number} Distance (mi)") * METERS_PER_MILE
            )
            seconds = round(
                positive_number(row[4], f"Row {number} Duration (min)") * 60
            )
            if seconds <= 0:
                raise ValueError(f"Row {number}: duration rounds to zero seconds")
            steps = row[6]
            if steps is not None and (
                isinstance(steps, bool)
                or not isinstance(steps, (int, float))
                or not math.isfinite(steps)
                or steps < 0
                or int(steps) != steps
            ):
                raise ValueError(
                    f"Row {number}: Running steps must be a nonnegative integer or blank"
                )
            category = text(row[2]) or "Unspecified"
            activities.append(
                {
                    "run_id": run_id,
                    "record_id": record_id,
                    "name": f"Treadmill · {category}",
                    "type": "Run",
                    "subtype": category,
                    "distance": round(distance, 6),
                    "moving_time": f"{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}",
                    "start_date_local": when.isoformat(timespec="seconds"),
                    "start_date": when.replace(tzinfo=ZoneInfo("America/New_York"))
                    .astimezone(UTC)
                    .isoformat()
                    .replace("+00:00", "Z"),
                    "date_precision": (
                        "day" if when.time() == datetime.min.time() else "minute"
                    ),
                    "location_country": None,
                    "summary_polyline": None,
                    "average_heartrate": None,
                    "elevation_gain": None,
                    "average_speed": distance / seconds,
                    "source": "manual_excel_treadmill",
                    "streak": 1,
                    "running_steps": int(steps) if steps is not None else None,
                    "workout_details": text(row[5]) or None,
                    "notes": text(row[7]) or None,
                }
            )
        if not activities:
            raise ValueError(
                "No completed Actual training rows; preserving the last valid dataset"
            )
        activities.sort(
            key=lambda a: (a["start_date_local"], a["run_id"]), reverse=True
        )
        check_ins = []
        if "Weekly Check-in" in workbook.sheetnames:
            weeks = set()
            for number, row in enumerate(
                workbook["Weekly Check-in"].iter_rows(values_only=True), 1
            ):
                if len(row) < 5 or text(row[4]).lower() != "actual" or row[0] is None:
                    continue
                start = recorded_date(row[0], f"Weekly Check-in row {number}").date()
                status = text(row[2])
                if (
                    start.weekday() != 0
                    or start in weeks
                    or status not in {"Complete", "In progress", "Needs update"}
                ):
                    raise ValueError(
                        f"Weekly Check-in row {number}: invalid/duplicate Monday or status"
                    )
                weeks.add(start)
                check_ins.append(
                    {
                        "week_start": start.isoformat(),
                        "week_end": (start + timedelta(days=6)).isoformat(),
                        "status": status,
                        "notes": text(row[3]) or None,
                    }
                )
        manifest = {
            "version": "1.1",
            "input_file": path.name,
            "input_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "timezone": "America/New_York",
            "input_distance_unit": "mi",
            "internal_distance_unit": "m",
            "display_distance_unit": "mi",
            "record_count": len(activities),
            "distance_miles": round(
                sum(a["distance"] for a in activities) / METERS_PER_MILE, 6
            ),
            "duration_seconds": sum(
                round(a["distance"] / a["average_speed"]) for a in activities
            ),
            "weekly_check_ins": sorted(check_ins, key=lambda a: a["week_start"]),
        }
        return activities, manifest
    finally:
        workbook.close()


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(
        dir=path.parent, prefix=path.name, suffix=".tmp"
    )
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as output:
            json.dump(value, output, ensure_ascii=False, indent=2, allow_nan=False)
            output.write("\n")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output", type=Path, default=Path("src/static/activities.json")
    )
    parser.add_argument(
        "--manifest", type=Path, default=Path("src/static/data-source.json")
    )
    args = parser.parse_args()
    try:
        activities, manifest = import_workbook(args.workbook)
    except (ValueError, OSError, KeyError) as error:
        print(f"Workbook validation failed: {error}", file=sys.stderr)
        return 1
    write_json(args.output, activities)
    write_json(args.manifest, manifest)
    print(
        f"Imported {len(activities)} Actual runs: {manifest['distance_miles']:g} mi, {manifest['duration_seconds'] // 60} min"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
